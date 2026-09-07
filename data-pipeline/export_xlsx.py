import json, csv, re, os
from collections import Counter, defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

rows=json.load(open('rows.json'))
COLS=list(rows[0].keys())

# ---- nettoyage léger ----
def clean_phone(p):
    if not p or p=="Non trouvé": return p
    p=re.sub(r'\s+',' ',p).strip()
    return p[:40]
for r in rows:
    r["Téléphone"]=clean_phone(r["Téléphone"])
    r["Nom du producteur"]=re.sub(r'\s+',' ',r["Nom du producteur"]).strip()

# ---- tri : Pays puis catégorie puis nom ----
def cat_of(r):
    p=(r["Produits principaux"] or "").lower()
    def has(*kw): return any(k in p for k in kw)
    if has("café","coffee"): return "Café"
    if has("cacao","cocoa","chocolat"): return "Cacao"
    if has("thé","tea") and not has("herbal"): return "Thé"
    if has("miel","honey","bee","apicole","cire"): return "Miel & produits de la ruche"
    if has("argan","olive","huile","oil") and not has("essential"): return "Huiles"
    if has("huile essentielle","essential oil","plante médicinale","medicinal","aromatique","infusion","herbal","herbes","plantes"): return "Plantes, herbes & huiles essentielles"
    if has("épice","spice","poivre","pepper","cardamome","gingembre","ginger","curcuma","turmeric","cannelle","vanille","vanilla","clou de girofle","cumin"): return "Épices"
    if has("karité","shea","cajou","cashew","noix","nut","amande","almond","sésame","sesame","oléagineux","oilseed"): return "Fruits à coque & oléagineux"
    if has("banane","banana","mangue","mango","ananas","pineapple","fruit","agrume","citrus","citron","orange","raisin","grape","avocat","avocado","papaye","pomme","poire","fraise","melon"): return "Fruits"
    if has("légume","vegetable","maraîch","horticult","tomate","oignon","chou","carotte","pomme de terre","potato","onion","okra","aubergine","brinjal"): return "Légumes & maraîchage"
    if has("coton","cotton","textile","fashion","vêtement","garment","tissu","artisanat","craft","décor","home","stationery","papier","jouet","bijou","jewell","ceramic","poterie","vannerie","basket"): return "Textile, artisanat & décoration"
    if has("cosmétique","beauty","bien-être","savon","soap","baume","crème"): return "Cosmétiques naturels"
    if has("sucre","sugar","panela","sirop"): return "Sucre / panela"
    if has("fleur","flower","rose","plante ornementale"): return "Fleurs & plantes"
    if has("vin","wine","raisin de cuve"): return "Vin"
    if has("quinoa","céréale","cereal","riz","rice","paddy","maïs","maize","blé","wheat","orge","barley","mil","millet","sorgho","avoine","kiwicha","chia","lin","linseed"): return "Céréales & graines"
    if has("lait","dairy","fromage","cheese","œuf","egg","viande","meat","volaille","poultry","poisson","fish","crevette"): return "Élevage & produits animaliers"
    if has("transformation","processed","jus","juice","confiture","jam","chips","farine","flour","séché","dried","torréfié","roasted","bière","beer","cacao en poudre"): return "Produits transformés"
    if p in ("non trouvé","non précisé",""): return "Non précisé"
    return "Autres produits agricoles"
for r in rows: r["_cat"]=cat_of(r)
rows.sort(key=lambda r:(r["Pays"] or "zz", r["_cat"], r["Nom du producteur"]))
for i,r in enumerate(rows,1): r["ID"]=f"ETH-{i:05d}"

with open("base_producteurs_ethimarket.csv","w",newline="",encoding="utf-8-sig") as f:
    w=csv.DictWriter(f,fieldnames=COLS,extrasaction="ignore"); w.writeheader()
    for r in rows: w.writerow({c:r.get(c,"") for c in COLS})

# ---- XLSX ----
wb=openpyxl.Workbook()
ws=wb.active; ws.title="Producteurs"
hdrfill=PatternFill("solid",fgColor="1F6B4A"); hdrfont=Font(color="FFFFFF",bold=True,size=10)
thin=Border(bottom=Side(style="thin",color="DDDDDD"))
ws.append(COLS)
for c in range(1,len(COLS)+1):
    cell=ws.cell(1,c); cell.fill=hdrfill; cell.font=hdrfont; cell.alignment=Alignment(vertical="center",wrap_text=True)
for r in rows:
    ws.append([r.get(c,"") for c in COLS])
ws.freeze_panes="C2"
ws.auto_filter.ref=f"A1:{get_column_letter(len(COLS))}{len(rows)+1}"
widths={"ID":10,"Nom du producteur":46,"Pays":16,"Continent":16,"Région / Ville":20,"Type de producteur":32,
 "Produits principaux":26,"Bio":8,"Équitable":9,"Certification":30,"Site internet":28,"Email":30,"Téléphone":22,
 "WhatsApp":18,"LinkedIn":22,"Instagram":22,"Facebook":22,"Contact / Responsable":24,"Adresse":34,"Source":40,"Notes":34}
for i,c in enumerate(COLS,1):
    ws.column_dimensions[get_column_letter(i)].width=widths.get(c,18)
for row in ws.iter_rows(min_row=2,max_row=len(rows)+1):
    for cell in row:
        cell.alignment=Alignment(vertical="top",wrap_text=False); cell.border=thin; cell.font=Font(size=9)

# --- Synthèse ---
w2=wb.create_sheet("Synthèse")
def put(row,vals,bold=False,fill=None):
    for i,v in enumerate(vals,1):
        c=w2.cell(row,i,v)
        if bold: c.font=Font(bold=True,color="FFFFFF" if fill else "000000",size=10)
        if fill: c.fill=PatternFill("solid",fgColor=fill)
    return row+1
r=1
r=put(r,["BASE MONDIALE DE PRODUCTEURS — ETHIMARKET"],bold=True,fill="1F6B4A")
r=put(r,[f"Générée le {__import__('datetime').date.today().isoformat()} — sources publiques uniquement, aucune donnée inventée"])
r+=1
r=put(r,["INDICATEUR","VALEUR"],bold=True,fill="1F6B4A")
def cnt(f): return sum(1 for x in rows if f(x))
ok=lambda v: v not in ("Non trouvé","","À vérifier")
stats=[("Nombre total de producteurs",len(rows)),
 ("Nombre de pays couverts",len(set(x["Pays"] for x in rows if x["Pays"]))),
 ("Producteurs avec email",cnt(lambda x: ok(x["Email"]))),
 ("Producteurs avec téléphone",cnt(lambda x: ok(x["Téléphone"]))),
 ("Producteurs avec site internet",cnt(lambda x: ok(x["Site internet"]))),
 ("Producteurs avec WhatsApp",cnt(lambda x: ok(x["WhatsApp"]))),
 ("Producteurs avec contact/responsable identifié",cnt(lambda x: ok(x["Contact / Responsable"]))),
 ("Producteurs avec LinkedIn",cnt(lambda x: ok(x["LinkedIn"]))),
 ("Producteurs avec Instagram",cnt(lambda x: ok(x["Instagram"]))),
 ("Producteurs avec Facebook",cnt(lambda x: ok(x["Facebook"]))),
 ("Producteurs certifiés commerce équitable (Fairtrade/FLOCERT ou WFTO)",cnt(lambda x: x["Équitable"]=="Oui")),
]
for k,v in stats: r=put(r,[k,v])
r+=1
r=put(r,["PRINCIPALES CATÉGORIES DE PRODUITS","NOMBRE"],bold=True,fill="1F6B4A")
for k,v in Counter(x["_cat"] for x in rows).most_common(): r=put(r,[k,v])
r+=1
r=put(r,["RÉPARTITION PAR CONTINENT","PRODUCTEURS","DONT AVEC EMAIL"],bold=True,fill="1F6B4A")
cont=defaultdict(list)
for x in rows: cont[x["Continent"] or "Non précisé"].append(x)
for k in sorted(cont,key=lambda k:-len(cont[k])):
    r=put(r,[k,len(cont[k]),sum(1 for x in cont[k] if ok(x["Email"]))])
r+=1
r=put(r,["PAYS","PRODUCTEURS","AVEC EMAIL","AVEC TÉLÉPHONE","AVEC SITE"],bold=True,fill="1F6B4A")
byp=defaultdict(list)
for x in rows: byp[x["Pays"] or "Non précisé"].append(x)
for k in sorted(byp,key=lambda k:-len(byp[k])):
    g=byp[k]
    r=put(r,[k,len(g),sum(1 for x in g if ok(x["Email"])),sum(1 for x in g if ok(x["Téléphone"])),sum(1 for x in g if ok(x["Site internet"]))])
for col,wd in zip("ABCDE",[62,14,14,16,12]): w2.column_dimensions[col].width=wd

# --- Sources & méthode ---
w3=wb.create_sheet("Sources & méthode")
lines=[
 ["SOURCE","TYPE","PRODUCTEURS","CE QUI A ÉTÉ COLLECTÉ","FIABILITÉ / LIMITES"],
 ["Fairtrade Finder — FLOCERT / Fairtrade International (API officielle assurance-ix)","Base officielle de certification",cnt(lambda x:"FLOCERT" in x["Source"]),"Nom, FLO-ID, adresse, ville, pays, produits certifiés, standard (SPO / Hired Labour), statut de certification","Source officielle et à jour. Pas d'email/téléphone publiés dans l'annuaire → champs contact à compléter (voir plan d'enrichissement)."],
 ["WFTO — World Fair Trade Organization (profils membres vérifiés)","Annuaire d'entreprises de commerce équitable",cnt(lambda x:"WFTO" in x["Source"]),"Nom, pays, ville, catégories produits, site, email, mobile, fixe, WhatsApp, réseaux sociaux","Données déclaratives des membres, vérifiées par WFTO."],
 ["Annuaire Clima y Café (Colombie) — coopératives et associations de caficulteurs","Annuaire sectoriel",cnt(lambda x:"climaycafe" in x["Source"]),"Nom, sigle, département, commune, adresse, téléphone, emails, représentant","Données compilées à partir de sources officielles (MinTIC / FNC). Annuaire non daté : téléphones à revalider."],
 ["PromPerú — Directorio de Productores y Exportadores de Productos Orgánicos del Perú","Annuaire officiel export",cnt(lambda x:"PromPerú" in x["Source"]),"Raison sociale, contact, 2 adresses, téléphone, fax, emails, site, produits","Officiel. Certains contacts peuvent avoir changé."],
 ["NSTIAM / NABARD — Directory of Farmer Producer Companies (Inde)","Annuaire officiel",cnt(lambda x:"NSTIAM" in x["Source"]),"Nom de la FPC, adresse, contact + fonction, mobile, email officiel, site, commodities, nombre de membres","Officiel (506 pages). FPC = sociétés de producteurs détenues par les agriculteurs."],
 ["Spices Board India — State wise FPO/FPC details","Annuaire officiel",cnt(lambda x:"Spices Board" in x["Source"]),"Nom, État, téléphones, emails, épices produites","Extraction partielle (mise en page tabulaire complexe) — noms à revalider."],
 ["TNAU Agritech — Organic Exporters (Inde)","Annuaire universitaire",cnt(lambda x:"TNAU" in x["Source"]),"Nom, adresse, État, téléphone, email, produits","Liste d'exportateurs bio (thé, café, épices) ; certains sont des planteurs-transformateurs."],
 ["Maroc Annuaire — catégorie Huile d'argane","Annuaire professionnel",cnt(lambda x:"Maroc Annuaire" in x["Source"]),"Nom, ville, adresse, téléphone, email, site","30 premières fiches de la catégorie (coopératives d'argan, dont coopératives féminines)."],
 ["Conseil oléicole international — exportateurs/importateurs huiles d'olive (Maroc)","Liste officielle internationale",cnt(lambda x:"Conseil oléicole" in x["Source"]),"Entité, adresse, téléphone, fax, emails","Mélange de coopératives, huileries et industriels — à filtrer selon vos critères."],
 ["","","","",""],
 ["RÈGLE APPLIQUÉE","","","","Aucune information inventée. Toute donnée absente = « Non trouvé ». Donnée incertaine = « À vérifier »."],
 ["FILTRES ANTI-REVENDEURS","","","","Les bases Fairtrade (standard « Producer » uniquement) et FPC/FPO/coopératives garantissent des producteurs réels. Les listes contenant des traders purs (exportateurs de café Ouganda, annuaire Ecocert) ont été écartées ou signalées."],
 ["À VENIR (passe 2)","","","","Enrichissement email/site/téléphone des 1 732 organisations de producteurs Fairtrade via leurs sites officiels ; annuaires bio nationaux (Agence Bio, SENASA, PGS India) ; coopératives d'Afrique de l'Ouest (cacao, karité, anacarde) ; artisans Océanie & Asie du Sud-Est."],
]
for i,l in enumerate(lines,1):
    for j,v in enumerate(l,1):
        c=w3.cell(i,j,v)
        if i==1: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F6B4A")
        c.alignment=Alignment(vertical="top",wrap_text=True)
for col,wd in zip("ABCDE",[46,26,12,60,60]): w3.column_dimensions[col].width=wd

wb.save("base_producteurs_ethimarket.xlsx")
print("OK lignes:",len(rows),"| pays:",len(set(x['Pays'] for x in rows)))
print("emails:",cnt(lambda x: ok(x["Email"])),"| tel:",cnt(lambda x: ok(x["Téléphone"])),"| sites:",cnt(lambda x: ok(x["Site internet"])))
print(Counter(x["_cat"] for x in rows).most_common(8))
